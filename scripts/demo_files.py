#!/usr/bin/env python3
"""Writes the demo folder the README clips and screenshots are made from: a
school project about the Space Shuttle ("Shuttle project") with notes, a
fact sheet PDF, a spreadsheet, a CSV of landmark flights, a mission-profile
infographic (SVG), a report, NASA photos and footage, a piano
recording, an icon and a 3D model to print.

    python3 scripts/demo_files.py OUT_DIR [--chromium PATH]

The pictures, the video, the music and the model are real and freely
licensed (see MEDIA below); they are downloaded once into a cache
($XDG_CACHE_HOME/monoform-demo), checked against SHA-256 sums, and trimmed
here. The documents are written here. Needs Pillow, openpyxl, python-docx,
ffmpeg and a Chromium headless shell (for the PDF).
"""
import argparse
import glob
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import urllib.request
import zipfile

from PIL import Image

NASA = 'https://images-assets.nasa.gov/'

# name → (url, sha256, credit). All public domain or CC0, except the icon (MIT).
MEDIA = {
    'atlantis.jpg': (NASA + 'image/iss015e12854/iss015e12854~large.jpg',
                     'eba42825e564870b873f57472c2b87da771857243f28afc6d2574a3235f26126',
                     'Space Shuttle Atlantis seen from the ISS, STS-117, 10 June 2007 — NASA, ISS015-E-12854 (public domain)'),
    'discovery.jpg': (NASA + 'image/iss018e043308/iss018e043308~orig.jpg',
                      '4b3940abc2a42425f009d15b2597568834ce37ffb5d733f747adfbe79c90088f',
                      'Space Shuttle Discovery after undocking from the ISS, STS-119, 25 March 2009 — NASA, ISS018-E-043308 (public domain)'),
    'launch.jpg': (NASA + 'image/jsc2011e067588/jsc2011e067588~large.jpg',
                   '8cce5cfb1ab20af8d4cfa67ec8714caaac5672da782870b0655aab52defb45e5',
                   'Atlantis lifts off on STS-135, 8 July 2011 — NASA, JSC2011-E-067588 (public domain)'),
    'landing.mp4': (NASA + 'video/ksc_112709_sts129_landing1/ksc_112709_sts129_landing1~orig.mp4',
                    'dcd4a28ec898f292d643dc95cc8cc5224ad6459daac629ca939137220eea99cc',
                    'Atlantis lands at Kennedy Space Center after STS-129, 27 November 2009 — NASA, KSC-09-S-00143 (public domain)'),
    'goldberg.ogg': ('https://archive.org/download/OpenGoldbergVariations/Kimiko%20Ishizaka%20-%20J.S.%20Bach-%20-Open-%20Goldberg'
                     '%20Variations%2C%20BWV%20988%20%28Piano%29%20-%2002%20Variatio%201%20a%201%20Clav..ogg',
                     '27901978812880ae5ce674089dd60aa3f6a0f33d1abd9b8186af0eac55b04fd1',
                     'J. S. Bach, Goldberg Variations, Variation 1 — Kimiko Ishizaka, The Open Goldberg Variations (CC0)'),
    'benchy.stl': ('https://raw.githubusercontent.com/CreativeTools/3DBenchy/b542785f9642ab39db8a778a0ccc934b3ebd829a/Single-part/3DBenchy.stl',
                   '6ab57f1c3f8e86bc3cbd302c6fa6270acf06277c6335454e922419c25d42e97e',
                   '#3DBenchy by Creative Tools (Daniel Norée, Paulo Kiefe) — public domain (CC0) since 2025'),
    'rocket.svg': ('https://raw.githubusercontent.com/microsoft/fluentui-emoji/1ffb34c752ecf5d402f04cfb4b392c77f57c54bc/assets/Rocket/Color/rocket_color.svg',
                   '3e759cbfacd8b9fb06df57abaf472bb2dbc1415bd60f21dc9f430e20fa0ca252',
                   'Rocket from Fluent Emoji — Microsoft (MIT)'),
}


def fetch(name):
    """The cached original of a MEDIA entry, downloaded once and verified."""
    url, sha, _ = MEDIA[name]
    cache = os.path.join(os.environ.get('XDG_CACHE_HOME') or os.path.expanduser('~/.cache'), 'monoform-demo')
    os.makedirs(cache, exist_ok=True)
    path = os.path.join(cache, name)
    if not os.path.exists(path) or hashlib.sha256(open(path, 'rb').read()).hexdigest() != sha:
        req = urllib.request.Request(url, headers={'User-Agent': 'Monoform-README/1.0 (https://github.com/caffeine-runtime/monoform)'})
        with urllib.request.urlopen(req, timeout=300) as r, open(path + '.part', 'wb') as f:
            shutil.copyfileobj(r, f)
        got = hashlib.sha256(open(path + '.part', 'rb').read()).hexdigest()
        if got != sha:
            raise SystemExit(f'{name}: checksum {got}, expected {sha}')
        os.replace(path + '.part', path)
    return path


# ------------------------------------------------------------------ notes

NOTES = """# Space Shuttle — project notes

![Atlantis lifts off on STS-135, the last flight](<STS-135 launch.jpg>)

Five orbiters flew **135 missions** between 1981 and 2011. These notes are
for the poster and the short talk on Friday.

## A mission, start to finish

```mermaid
flowchart LR
  A[Launch] --> B[Orbit]
  B --> C[Dock with the ISS]
  C --> D[Re-entry]
  D --> E[Runway landing]
```

## The fleet

| Orbiter | Flights | In service |
|---------|--------:|------------|
| Columbia | 28 | 1981–2003 |
| Challenger | 10 | 1983–1986 |
| Discovery | 39 | 1984–2011 |
| Atlantis | 33 | 1985–2011 |
| Endeavour | 25 | 1992–2011 |

## To do

- ~~find launch photos~~ (NASA's are public domain)
- print a Benchy to test the printer before the model
- ask about the projector

> The last landing: Atlantis, 21 July 2011, Kennedy Space Center.
"""

# An older draft of the same notes, to compare with.
DRAFT = (NOTES.replace('Five orbiters flew **135 missions** between 1981 and 2011.', 'The orbiters flew more than a hundred missions.')
         .replace('  C --> D[Re-entry]\n', '  B --> D[Re-entry]\n').replace('| Endeavour | 25 | 1992–2011 |\n', '')
         .replace('- ask about the projector\n', '- ask about the projector\n- find a video of a landing\n'))

POSTER = """# Poster

- the launch photo across the top
- the fleet table
- a printed Benchy on the desk
"""

MOODBOARD = """# Moodboard

Cut-outs for the poster:

"""

# ---------------------------------------------------------- the fact sheet

SHEET = """<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 0; }}
body {{ margin: 0; font-family: 'Source Serif 4', Georgia, serif; color: #1d2433; }}
.page {{ width: 210mm; height: 297mm; box-sizing: border-box; padding: 18mm 20mm; page-break-after: always; position: relative; }}
.cover {{ padding: 0; background: #0f1524; color: #eef2f8; }}
.cover img {{ width: 100%; height: 185mm; object-fit: cover; display: block; }}
.cover .t {{ padding: 14mm 20mm; }}
h1 {{ font-size: 46pt; margin: 0 0 4mm; font-weight: 600; letter-spacing: -0.5pt; }}
h2 {{ font-size: 20pt; margin: 0 0 4mm; color: #1f4e96; }}
p, li {{ font-size: 11.5pt; line-height: 1.55; }}
.kicker {{ font: 600 10pt Inter, sans-serif; letter-spacing: 2pt; text-transform: uppercase; color: #e0a33a; }}
.cols {{ columns: 2; column-gap: 10mm; }}
.box {{ background: #eef2f8; border-radius: 4mm; padding: 6mm 8mm; margin: 6mm 0; }}
.photo {{ width: 100%; border-radius: 3mm; margin: 4mm 0; }}
table {{ border-collapse: collapse; width: 100%; font: 10.5pt Inter, sans-serif; }}
td, th {{ border-bottom: 1px solid #d5dbe6; padding: 2.5mm 2mm; text-align: left; }}
.foot {{ position: absolute; bottom: 12mm; left: 20mm; right: 20mm; font: 9pt Inter, sans-serif; color: #7a8394; display: flex; justify-content: space-between; }}
</style></head><body>
<div class="page cover"><img src="launch.jpg"><div class="t"><div class="kicker">Fact sheet · 1981–2011</div>
<h1>The Space Shuttle</h1><p>Five orbiters, 135 missions and thirty years of flying to orbit and back on wings.</p></div></div>
<div class="page"><div class="kicker">How a mission worked</div><h2>Up like a rocket, down like a glider</h2><div class="cols">
<p>The orbiter launched upright, strapped to a large external fuel tank and two solid rocket boosters. The boosters dropped
away about two minutes after liftoff and were recovered from the ocean; the tank was released just before orbit.</p>
<p>In orbit the payload bay doors opened. The bay was about 18 metres long — big enough for satellites, laboratories and
the modules that built the International Space Station.</p>
<p>To come home, the orbiter fired its engines against the direction of travel, fell back into the atmosphere protected
by its heat-shield tiles, and landed on a runway without power.</p></div>
<img class="photo" src="atlantis.jpg">
<div class="box"><b>The back-flip.</b> Before docking with the ISS, each orbiter turned a slow somersault so that the
station crew could photograph its heat shield — the photo above is one of those pictures.</div>
<div class="foot"><span>The Space Shuttle</span><span>2</span></div></div>
<div class="page"><div class="kicker">The fleet</div><h2>Five orbiters</h2>
<table><tr><th>Orbiter</th><th>First flight</th><th>Last flight</th><th>Flights</th></tr>
<tr><td>Columbia</td><td>1981</td><td>2003</td><td>28</td></tr><tr><td>Challenger</td><td>1983</td><td>1986</td><td>10</td></tr>
<tr><td>Discovery</td><td>1984</td><td>2011</td><td>39</td></tr><tr><td>Atlantis</td><td>1985</td><td>2011</td><td>33</td></tr>
<tr><td>Endeavour</td><td>1992</td><td>2011</td><td>25</td></tr><tr><th>Total</th><th></th><th></th><th>135</th></tr></table>
<div class="box"><b>Where they are now.</b> Discovery is at the Steven F. Udvar-Hazy Center in Virginia, Atlantis at the
Kennedy Space Center Visitor Complex in Florida and Endeavour at the California Science Center in Los Angeles.</div>
<p style="font-size:9pt;color:#7a8394">Photos: NASA (public domain).</p>
<div class="foot"><span>The Space Shuttle</span><span>3</span></div></div>
</body></html>"""


def fact_sheet(path, launch, atlantis, chromium):
    with tempfile.TemporaryDirectory() as tmp:
        shutil.copy(launch, os.path.join(tmp, 'launch.jpg'))
        shutil.copy(atlantis, os.path.join(tmp, 'atlantis.jpg'))
        html = os.path.join(tmp, 'sheet.html')
        with open(html, 'w') as f:
            f.write(SHEET.format())
        subprocess.run([chromium, '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={path}', 'file://' + html],
                       check=True, capture_output=True)


# ------------------------------------------------------------ spreadsheet

FLEET = [('Columbia', 1981, 2003, 28), ('Challenger', 1983, 1986, 10), ('Discovery', 1984, 2011, 39),
         ('Atlantis', 1985, 2011, 33), ('Endeavour', 1992, 2011, 25)]


def fleet_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Fleet'
    ws.append(['Orbiter', 'First flight', 'Last flight', 'Flights', 'Years in service', 'Share of flights'])
    for i, (name, first, last, flights) in enumerate(FLEET, 2):
        ws.append([name, first, last, flights, f'=C{i}-B{i}', f'=ROUND(D{i}/D$7*100,1)'])
    ws.append(['Total', '=MIN(B2:B6)', '=MAX(C2:C6)', '=SUM(D2:D6)', '=C7-B7'])
    for c in ws[1] + ws[7]:
        c.font = Font(bold=True)
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='EEF2F8')
    for col, w in zip('ABCDEF', (14, 12, 12, 10, 16, 16)):
        ws.column_dimensions[col].width = w
    wb.save(path)


# ------------------------------------------------------------------- CSV

# Landmark flights: (mission, orbiter, launch, landing, highlight). Days are worked out.
FLIGHTS = [
    ('STS-1', 'Columbia', '1981-04-12', '1981-04-14', 'First orbital flight'),
    ('STS-7', 'Challenger', '1983-06-18', '1983-06-24', 'Sally Ride, first American woman in space'),
    ('STS-9', 'Columbia', '1983-11-28', '1983-12-08', 'First Spacelab'),
    ('STS-41-B', 'Challenger', '1984-02-03', '1984-02-11', 'First untethered spacewalk'),
    ('STS-41-D', 'Discovery', '1984-08-30', '1984-09-05', "Discovery's first flight"),
    ('STS-51-J', 'Atlantis', '1985-10-03', '1985-10-07', "Atlantis's first flight"),
    ('STS-26', 'Discovery', '1988-09-29', '1988-10-03', 'Return to flight'),
    ('STS-31', 'Discovery', '1990-04-24', '1990-04-29', 'Hubble Space Telescope deployed'),
    ('STS-49', 'Endeavour', '1992-05-07', '1992-05-16', "Endeavour's first flight"),
    ('STS-61', 'Endeavour', '1993-12-02', '1993-12-13', 'First Hubble servicing'),
    ('STS-71', 'Atlantis', '1995-06-27', '1995-07-07', 'First docking with Mir'),
    ('STS-88', 'Endeavour', '1998-12-04', '1998-12-15', 'First ISS assembly flight (Unity)'),
    ('STS-114', 'Discovery', '2005-07-26', '2005-08-09', 'Return to flight'),
    ('STS-117', 'Atlantis', '2007-06-08', '2007-06-22', 'The photo on the fact sheet'),
    ('STS-119', 'Discovery', '2009-03-15', '2009-03-28', 'The photo in this folder'),
    ('STS-125', 'Atlantis', '2009-05-11', '2009-05-24', 'Last Hubble servicing'),
    ('STS-129', 'Atlantis', '2009-11-16', '2009-11-27', 'The landing video'),
    ('STS-133', 'Discovery', '2011-02-24', '2011-03-09', "Discovery's last flight"),
    ('STS-134', 'Endeavour', '2011-05-16', '2011-06-01', "Endeavour's last flight"),
    ('STS-135', 'Atlantis', '2011-07-08', '2011-07-21', 'Last flight of the programme'),
]


def flights_csv(path):
    import csv
    from datetime import date

    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Mission', 'Orbiter', 'Launch', 'Landing', 'Days', 'Highlight'])
        for m, o, a, b, h in FLIGHTS:
            w.writerow([m, o, a, b, (date.fromisoformat(b) - date.fromisoformat(a)).days, h])


# ------------------------------------------------------------------- SVG

def mission_svg(path):
    """A launch-to-landing infographic: a curved Earth, altitude rings, the
    flight path with its eight phases, the ISS and the orbiter (schematic)."""
    import math
    import random

    W, H = 1600, 1000
    cx, cy, R = 800, 2330, 1640  # Earth's centre and radius in the drawing; its top is at y = 690
    k = 1.0  # pixels per km of altitude
    rnd = random.Random(1981)
    out = []
    a = out.append

    def pt(deg, alt):
        t = math.radians(deg)
        r = R + alt * k
        return cx + r * math.sin(t), cy - r * math.cos(t)

    def poly(points):
        return 'M' + ' L'.join(f'{x:.1f},{y:.1f}' for x, y in points)

    a(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" width="{W}" height="{H}" font-family="Inter, Segoe UI, Helvetica, Arial, sans-serif">')
    a('<title>Space Shuttle mission profile</title>')
    a('<defs>')
    a('<linearGradient id="sky" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#03050d"/><stop offset=".55" stop-color="#0a1430"/><stop offset="1" stop-color="#17305f"/></linearGradient>')
    a(f'<radialGradient id="earth" cx="{cx}" cy="{cy - R * 0.55:.0f}" r="{R * 1.1:.0f}" gradientUnits="userSpaceOnUse"><stop offset="0" stop-color="#3f8fd8"/><stop offset=".45" stop-color="#1d5aa6"/><stop offset="1" stop-color="#0a234d"/></radialGradient>')
    a(f'<radialGradient id="air" cx="{cx}" cy="{cy}" r="{R + 70}" gradientUnits="userSpaceOnUse"><stop offset="{(R - 10) / (R + 70):.4f}" stop-color="#7cc6ff" stop-opacity=".75"/><stop offset="{(R + 20) / (R + 70):.4f}" stop-color="#4aa3ff" stop-opacity=".28"/><stop offset="1" stop-color="#4aa3ff" stop-opacity="0"/></radialGradient>')
    a('<linearGradient id="up" x1="0" y1="1" x2="1" y2="0"><stop offset="0" stop-color="#ffd166"/><stop offset="1" stop-color="#ff7a3d"/></linearGradient>')
    a('<linearGradient id="down" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#ff7aa8"/><stop offset="1" stop-color="#ff4d5e"/></linearGradient>')
    a('<linearGradient id="panel" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#f4c95d"/><stop offset="1" stop-color="#b8862b"/></linearGradient>')
    a('<linearGradient id="card" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#16213f" stop-opacity=".92"/><stop offset="1" stop-color="#0d1530" stop-opacity=".92"/></linearGradient>')
    a(f'<clipPath id="globe"><circle cx="{cx}" cy="{cy}" r="{R}"/></clipPath>')
    a('</defs>')
    a(f'<rect width="{W}" height="{H}" fill="url(#sky)"/>')
    # Stars, a few brighter ones with a soft halo.
    a('<g fill="#fff">')
    for _ in range(260):
        x, y = rnd.uniform(0, W), rnd.uniform(0, 640)
        r = rnd.choice((0.5, 0.6, 0.8, 1.0, 1.3)) if rnd.random() > 0.04 else 1.9
        a(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{r}" opacity="{rnd.uniform(0.35, 0.95):.2f}"/>')
        if r > 1.5:
            a(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="6" opacity=".12"/>')
    a('</g>')
    # Earth: atmosphere glow, the globe, latitude and longitude lines, city lights.
    a(f'<circle cx="{cx}" cy="{cy}" r="{R + 70}" fill="url(#air)"/>')
    a(f'<circle cx="{cx}" cy="{cy}" r="{R}" fill="url(#earth)"/>')
    a('<g clip-path="url(#globe)" fill="none" stroke="#9fd0ff" stroke-opacity=".14" stroke-width="1.2">')
    for i in range(1, 9):
        a(f'<ellipse cx="{cx}" cy="{cy}" rx="{R * i / 9:.0f}" ry="{R}"/>')
    for i in range(1, 12):
        a(f'<circle cx="{cx}" cy="{cy + R * 0.1}" r="{R * (1 - i * 0.012):.0f}"/>')
    a('</g>')
    a('<g fill="#ffd98a">')
    for _ in range(140):
        deg = rnd.uniform(-30, 30)
        x, y = pt(deg, -rnd.uniform(4, 120))
        a(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{rnd.choice((0.7, 1, 1.4))}" opacity="{rnd.uniform(0.25, 0.7):.2f}"/>')
    a('</g>')
    # Altitude rings.
    a('<g fill="none" stroke="#c9d8ff" stroke-opacity=".22" stroke-dasharray="4 8">')
    for alt in (100, 200, 300, 400):
        a(f'<circle cx="{cx}" cy="{cy}" r="{R + alt * k}"/>')
    a('</g>')

    # The flight: ascent, orbit, descent.
    asc = [(-24 + 13 * (i / 50) ** 0.75, 400 * (1 - math.exp(-4.2 * i / 50)) / (1 - math.exp(-4.2))) for i in range(51)]
    orb = [(-11 + 22 * i / 40, 400) for i in range(41)]
    des = [(11 + 13 * i / 50, 400 * (1 - (i / 50) ** 1.6) ** 1.25) for i in range(51)]
    for pts, stroke in ((asc, 'url(#up)'), (orb, '#6fe3ff'), (des, 'url(#down)')):
        line = poly(pt(d, h) for d, h in pts)
        a(f'<path d="{line}" fill="none" stroke="{stroke}" stroke-opacity=".25" stroke-width="12" stroke-linecap="round"/>')
        a(f'<path d="{line}" fill="none" stroke="{stroke}" stroke-width="3.5" stroke-linecap="round"/>')
    # Direction ticks along the orbit.
    for d in (-6, 0, 6):
        x, y = pt(d, 400)
        a(f'<path d="M{x - 7:.1f},{y - 6:.1f} L{x + 5:.1f},{y:.1f} L{x - 7:.1f},{y + 6:.1f}" fill="none" stroke="#6fe3ff" stroke-width="2.4" transform="rotate({d} {x:.1f} {y:.1f})"/>')

    # The ISS near the top of the orbit: truss, modules, eight solar wings, radiators.
    ix, iy = pt(3, 400)
    a(f'<g transform="translate({ix:.1f} {iy - 46:.1f}) rotate(3)">')
    a('<rect x="-62" y="-2.5" width="124" height="5" fill="#cfd6e4"/>')
    for sx in (-58, -42, 42, 58):
        for sy in (-1, 1):
            y0 = -30 if sy < 0 else 4
            a(f'<rect x="{sx - 6}" y="{y0}" width="12" height="26" fill="url(#panel)" stroke="#6b4e12" stroke-width=".6"/>')
            for gy in range(1, 4):
                a(f'<line x1="{sx - 6}" y1="{y0 + gy * 6.5}" x2="{sx + 6}" y2="{y0 + gy * 6.5}" stroke="#6b4e12" stroke-width=".5"/>')
    a('<rect x="-18" y="-6" width="36" height="12" rx="5" fill="#e9edf5"/><rect x="-6" y="-16" width="12" height="32" rx="4" fill="#dfe5ef"/>')
    a('<rect x="-30" y="-12" width="8" height="24" fill="#f5f7fb" opacity=".9"/><rect x="22" y="-12" width="8" height="24" fill="#f5f7fb" opacity=".9"/>')
    a('</g>')
    # The orbiter on its way to the station.
    ox, oy = pt(-6.5, 400)
    a(f'<g transform="translate({ox:.1f} {oy - 24:.1f}) rotate(-4) scale(1.15)">')
    a('<path d="M-30,0 C-26,-5 -10,-6 6,-6 L20,-17 L26,-17 L22,-6 L30,-4 C34,-2 34,2 30,4 L22,6 L26,17 L20,17 L6,6 C-10,6 -26,5 -30,0 Z" fill="#f2f4f8" stroke="#3a4560" stroke-width="1"/>')
    a('<path d="M-30,0 C-27,-3.5 -21,-4.5 -16,-4.5 L-16,4.5 C-21,4.5 -27,3.5 -30,0 Z" fill="#2b3245"/>')
    a('<path d="M24,-3 L36,-8 L36,8 L24,3 Z" fill="#ffb347" opacity=".85"/>')
    a('</g>')

    # Launch pad and runway.
    lx, ly = pt(-24, 0)
    a(f'<g transform="translate({lx:.1f} {ly:.1f}) rotate(-24)"><rect x="-3" y="-26" width="6" height="26" fill="#e8edf6"/><rect x="-11" y="-30" width="4" height="30" fill="#8794b3"/><rect x="-16" y="-2" width="32" height="4" fill="#8794b3"/></g>')
    rx, ry = pt(24, 0)
    a(f'<g transform="translate({rx:.1f} {ry:.1f}) rotate(24)"><rect x="-34" y="-3" width="68" height="5" fill="#5d6a88"/>'
      + ''.join(f'<rect x="{-30 + i * 9}" y="-1.2" width="5" height="1.4" fill="#fff"/>' for i in range(7)) + '</g>')

    # The eight phases: (angle, altitude, label, detail, label side).
    phases = [
        (-24, 0, 'Launch', 'Kennedy Space Center, Florida', 'below'),
        (-21.6, 150, 'Boosters separate', '≈ 2 min after launch · ≈ 45 km', 'right'),
        (-14.2, 360, 'Main engines cut off', '≈ 8½ min · the tank falls away', 'right-low'),
        (-9, 400, 'In orbit', '≈ 28 000 km/h · ≈ 350–400 km up', 'above'),
        (3, 400, 'Docking with the ISS', 'crew, cargo and new modules', 'above-r'),
        (11, 400, 'Deorbit burn', 'about an hour before landing', 'right'),
        (17.5, 140, 'Entry', 'heat shield up to ≈ 1 600 °C', 'left'),
        (24, 0, 'Landing', '≈ 350 km/h on the runway', 'below'),
    ]
    for n, (deg, alt, label, detail, side) in enumerate(phases, 1):
        x, y = pt(deg, alt)
        bw = max(len(label) * 10.5, len(detail) * 7.6) + 30
        tx, ty = {
            'left': (x - bw - 40, y - 22),
            'right': (x + 40, y - 22),
            'right-low': (x + 36, y + 8),
            'above': (x - 110, y - 120),
            'above-r': (x + 30, y - 150),
            'below': (x - 100, y + 52),
        }[side]
        anchor = 'start'
        a(f'<line x1="{x:.1f}" y1="{y:.1f}" x2="{tx + 14:.1f}" y2="{ty + 20:.1f}" stroke="#c9d8ff" stroke-opacity=".45" stroke-width="1.2"/>')
        a(f'<g transform="translate({tx:.1f} {ty:.1f})"><rect width="{bw:.0f}" height="56" rx="10" fill="url(#card)" stroke="#3a4f86" stroke-width="1"/>'
          f'<text x="15" y="24" font-size="17" font-weight="700" fill="#f3f6ff" text-anchor="{anchor}">{label}</text>'
          f'<text x="15" y="44" font-size="13" fill="#a9b8de" text-anchor="{anchor}">{detail}</text></g>')
        a(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="15" fill="#0d1530" stroke="#ffd166" stroke-width="2.5"/>'
          f'<text x="{x:.1f}" y="{y + 5.5:.1f}" font-size="15" font-weight="700" fill="#ffd166" text-anchor="middle">{n}</text>')

    # Title, legend and the honest small print.
    a('<g transform="translate(56 70)"><text font-size="15" font-weight="700" letter-spacing="4" fill="#ff9f5a">SPACE SHUTTLE · 1981–2011</text>'
      '<text y="52" font-size="46" font-weight="800" fill="#f3f6ff">Mission profile</text>'
      '<text y="84" font-size="18" fill="#a9b8de">From the launch pad to the runway: a typical flight took one to two weeks.</text></g>')
    a('<g transform="translate(1268 62)"><rect width="276" height="128" rx="12" fill="url(#card)" stroke="#3a4f86"/>')
    for i, (stroke, name) in enumerate((('#ffa94d', 'Ascent'), ('#6fe3ff', 'Orbit'), ('#ff5d73', 'Descent and entry'))):
        a(f'<line x1="22" y1="{32 + i * 32}" x2="70" y2="{32 + i * 32}" stroke="{stroke}" stroke-width="4" stroke-linecap="round"/>'
          f'<text x="86" y="{37 + i * 32}" font-size="15" fill="#e3e9fb">{name}</text>')
    a('</g>')
    a(f'<text x="{W - 40}" y="{H - 28}" font-size="13" fill="#7f8fb8" text-anchor="end">Schematic, not to scale.</text>')
    a('</svg>')
    with open(path, 'w') as f:
        f.write('\n'.join(out) + '\n')


# ----------------------------------------------------------------- report

def report_docx(path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading('The Space Shuttle', 0)
    p = doc.add_paragraph('A short report for the science fair. ')
    p.add_run('Sources: NASA history pages and press kits.').italic = True
    for title, text in (
        ('A reusable spacecraft', 'The Space Shuttle was the first spacecraft built to be flown again and again. The orbiter, '
         'its main engines and the two solid rocket boosters were reused; only the large external tank was new for every flight.'),
        ('What it carried', 'Satellites, laboratories and telescopes rode in the payload bay. Discovery carried the Hubble Space '
         'Telescope into orbit in April 1990, and many flights carried parts of the International Space Station.'),
        ('The last flight', 'Atlantis flew the final mission, STS-135, launching on 8 July 2011 and landing at Kennedy Space '
         'Center on 21 July 2011.'),
    ):
        doc.add_heading(title, level=2)
        doc.add_paragraph(text)
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Light Grid Accent 1'
    for c, h in zip(t.rows[0].cells, ('Orbiter', 'Flights', 'On display at')):
        c.text = h
    for r in (('Discovery', '39', 'Steven F. Udvar-Hazy Center, Virginia'), ('Atlantis', '33', 'Kennedy Space Center Visitor Complex, Florida'),
              ('Endeavour', '25', 'California Science Center, Los Angeles')):
        for c, v in zip(t.add_row().cells, r):
            c.text = v
    note = doc.add_paragraph().add_run('Draft — check the numbers against the fact sheet before printing.')
    note.font.size = Pt(9)
    note.font.color.rgb = RGBColor(0x7a, 0x83, 0x94)
    doc.save(path)


# ------------------------------------------------------- data and a page

MISSIONS = {
    'program': 'Space Shuttle',
    'years': [1981, 2011],
    'missions': [
        {'name': 'STS-1', 'orbiter': 'Columbia', 'launch': '1981-04-12', 'note': 'first flight'},
        {'name': 'STS-31', 'orbiter': 'Discovery', 'launch': '1990-04-24', 'note': 'carried the Hubble Space Telescope'},
        {'name': 'STS-88', 'orbiter': 'Endeavour', 'launch': '1998-12-04', 'note': 'first ISS assembly flight (Unity)'},
        {'name': 'STS-117', 'orbiter': 'Atlantis', 'launch': '2007-06-08', 'note': 'the photo on the fact sheet'},
        {'name': 'STS-119', 'orbiter': 'Discovery', 'launch': '2009-03-15', 'note': 'the photo in this folder'},
        {'name': 'STS-129', 'orbiter': 'Atlantis', 'landing': '2009-11-27', 'note': 'the landing video'},
        {'name': 'STS-135', 'orbiter': 'Atlantis', 'launch': '2011-07-08', 'landing': '2011-07-21', 'note': 'last flight'},
    ],
    'onDisplay': {'Discovery': 'Virginia', 'Atlantis': 'Florida', 'Endeavour': 'Los Angeles'},
}

PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Where the orbiters are now</title><style>
body { margin: 0; font-family: Georgia, serif; background: #eef2f8; color: #1d2433; }
header { padding: 48px 40px; background: #0f1524; color: #eef2f8; }
h1 { margin: 0; font-size: 40px; } header p { color: #e0a33a; letter-spacing: 2px; font: 600 12px sans-serif; text-transform: uppercase; }
main { padding: 24px 40px; display: grid; grid-template-columns: repeat(2, 1fr); gap: 16px; }
.card { background: #fff; border-radius: 12px; padding: 18px; box-shadow: 0 1px 3px #0002; }
.card h2 { margin: 0 0 6px; color: #1f4e96; font-size: 20px; }
</style></head><body><header><p>Field trip ideas</p><h1>Where the orbiters are now</h1></header><main>
<div class="card"><h2>Discovery</h2>Steven F. Udvar-Hazy Center, Chantilly, Virginia.</div>
<div class="card"><h2>Atlantis</h2>Kennedy Space Center Visitor Complex, Florida.</div>
<div class="card"><h2>Endeavour</h2>California Science Center, Los Angeles.</div>
<div class="card"><h2>Enterprise</h2>The test orbiter, at the Intrepid Museum in New York.</div>
</main></body></html>
"""


LANDING_START = 1.5


def find_chromium():
    base = os.environ.get('PLAYWRIGHT_BROWSERS_PATH', '/opt/pw-browsers')
    found = sorted(glob.glob(os.path.join(base, 'chromium_headless_shell-*', '*', 'headless_shell'))
                   + glob.glob(os.path.join(base, 'chromium_headless_shell-*', '*', 'chrome-headless-shell')))
    return found[-1] if found else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('out')
    ap.add_argument('--chromium', default=find_chromium())
    args = ap.parse_args()
    root = os.path.join(args.out, 'Shuttle project')
    shutil.rmtree(root, ignore_errors=True)
    os.makedirs(os.path.join(root, 'notes'))
    put = lambda name: os.path.join(root, name)  # noqa: E731

    # NASA's photos: Discovery over the Earth's horizon (the image editor's demo: a clear
    # outline against space and cloud tops) without the ID strip and the window's corner.
    Image.open(fetch('discovery.jpg')).crop((0, 150, 2000, 2350)).save(put('Discovery in orbit.jpg'), quality=92)
    shutil.copy(fetch('launch.jpg'), put('STS-135 launch.jpg'))
    # 22 seconds of the landing: gliding in, the runway, touchdown (after the fade-in).
    subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-ss', str(LANDING_START), '-i', fetch('landing.mp4'), '-t', '22', '-c:v', 'libvpx-vp9', '-b:v', '1.6M',
                    '-c:a', 'libopus', '-b:a', '96k', put('Atlantis landing.webm')], check=True)
    shutil.copy(fetch('goldberg.ogg'), put('Goldberg Variation 1.ogg'))
    shutil.copy(fetch('benchy.stl'), put('3DBenchy.stl'))
    shutil.copy(fetch('rocket.svg'), put('Rocket.svg'))

    for name, text in (('Notes.md', NOTES), ('notes/Notes (draft).md', DRAFT), ('notes/Poster.md', POSTER), ('notes/Moodboard.md', MOODBOARD)):
        with open(put(name), 'w') as f:
            f.write(text)
    # The fact sheet keeps Atlantis seen from the ISS (STS-117), without the ID strip along its bottom.
    with tempfile.TemporaryDirectory() as tmp:
        atlantis = os.path.join(tmp, 'atlantis.jpg')
        Image.open(fetch('atlantis.jpg')).crop((0, 0, 1920, 1248)).save(atlantis, quality=92)
        fact_sheet(put('Shuttle fact sheet.pdf'), put('STS-135 launch.jpg'), atlantis, args.chromium)
    fleet_xlsx(put('Fleet.xlsx'))
    flights_csv(put('Landmark flights.csv'))
    mission_svg(put('Mission profile.svg'))
    report_docx(put('Report.docx'))
    with open(put('missions.json'), 'w') as f:
        json.dump(MISSIONS, f, indent=2, ensure_ascii=False)
    with open(put('museums.html'), 'w') as f:
        f.write(PAGE)
    with zipfile.ZipFile(put('photos.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        for name in ('Discovery in orbit.jpg', 'STS-135 launch.jpg', 'Rocket.svg'):
            z.write(put(name), name)
    with open(put('CREDITS.txt'), 'w') as f:
        f.writelines(f'{credit}\n' for _, _, credit in MEDIA.values())
    sums = []
    for name in sorted(os.listdir(root)):
        if os.path.isfile(put(name)):
            sums.append(f'{hashlib.sha256(open(put(name), "rb").read()).hexdigest()}  {name}\n')
    with open(put('SHA256SUMS'), 'w') as f:
        f.writelines(sums)
    print(root)


if __name__ == '__main__':
    main()
